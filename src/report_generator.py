import json
from collections import Counter
from pathlib import Path
from typing import Any, Dict, List, Optional

import pandas as pd
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

# Confidence-level row fill colours (Excel hex, no leading #)
_FILL_HIGH = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")
_FILL_MED  = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
_FILL_LOW  = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
_FLAG_FILL = {"high": _FILL_HIGH, "medium": _FILL_MED, "low": _FILL_LOW}
_BOLD = Font(bold=True)


def save_results(
    results_df: pd.DataFrame,
    output_dir: str,
    base: str = "classification_results",
) -> Dict[str, str]:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    csv_path  = out / f"{base}.csv"
    xlsx_path = out / f"{base}.xlsx"

    results_df.to_csv(csv_path, index=False)

    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        results_df.to_excel(writer, sheet_name="Classification Results", index=False)
        ws = writer.sheets["Classification Results"]

        # --- Header styling ---
        for cell in ws[1]:
            cell.font = _BOLD
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions

        # --- Column widths ---
        for col_cells in ws.columns:
            max_len = max((len(str(cell.value or "")) for cell in col_cells), default=10)
            ws.column_dimensions[col_cells[0].column_letter].width = min(max_len + 2, 55)

        # --- Confidence-flag row colouring ---
        # Find the confidence_flag column index dynamically
        flag_col_idx = None
        for cell in ws[1]:
            if str(cell.value).lower() == "confidence_flag":
                flag_col_idx = cell.column
                break

        if flag_col_idx is not None:
            for row in ws.iter_rows(min_row=2, max_row=ws.max_row):
                flag_val = str(row[flag_col_idx - 1].value or "").lower()
                fill = _FLAG_FILL.get(flag_val)
                if fill:
                    for cell in row:
                        cell.fill = fill

    return {"csv": str(csv_path), "excel": str(xlsx_path)}


def generate_summary(
    results_df: pd.DataFrame,
    systems: List[str],
    data_quality: Optional[Dict[str, Any]] = None,
) -> Dict[str, Any]:
    summary: Dict[str, Any] = {
        "total_assets_processed": int(results_df["asset_id"].nunique()),
        "total_classification_records": len(results_df),
        "systems": {},
    }

    manual_review: set = set()
    top1 = results_df[results_df["match_rank"] == "1st"]

    for system in systems:
        sys_top1 = top1[top1["classification_system"] == system]
        if sys_top1.empty:
            continue
        counts = sys_top1["confidence_flag"].value_counts().to_dict()
        summary["systems"][system] = {
            "assets_classified": len(sys_top1),
            "average_top_score": round(float(sys_top1["similarity_score"].mean()), 1),
            "high_confidence": counts.get("high", 0),
            "medium_confidence": counts.get("medium", 0),
            "low_confidence": counts.get("low", 0),
            "flagged_for_review": counts.get("low", 0),
        }
        low_ids = sys_top1[sys_top1["confidence_flag"] == "low"]["asset_id"].tolist()
        manual_review.update(low_ids)

    summary["assets_requiring_manual_review"] = sorted(manual_review)
    summary["manual_review_count"] = len(manual_review)

    # --- Cross-system category consistency (O6) ---
    if "matched_category" in results_df.columns and len(systems) > 1:
        per_asset: Dict[str, bool] = {}
        for asset_id, grp in top1.groupby("asset_id"):
            cats = grp["matched_category"].dropna().tolist()
            if not cats:
                per_asset[str(asset_id)] = False
                continue
            most_common_count = Counter(cats).most_common(1)[0][1]
            # Majority (>50%) agree on same category → consistent
            per_asset[str(asset_id)] = most_common_count > len(systems) / 2

        consistent = sum(per_asset.values())
        total = len(per_asset)
        summary["cross_system_consistency"] = {
            "consistent_count": consistent,
            "inconsistent_count": total - consistent,
            "consistency_pct": round(consistent / total * 100, 1) if total else 0.0,
            "inconsistent_assets": [k for k, v in per_asset.items() if not v],
        }

    # --- Data quality (O7) ---
    if data_quality:
        summary["data_quality"] = data_quality

    return summary


def save_summary(
    summary: Dict[str, Any],
    output_dir: str,
    base: str = "classification_summary",
) -> Dict[str, str]:
    out = Path(output_dir)
    out.mkdir(parents=True, exist_ok=True)

    json_path = out / f"{base}.json"
    txt_path  = out / f"{base}.txt"

    with open(json_path, "w") as f:
        json.dump(summary, f, indent=2, default=str)

    lines = [
        "=" * 62,
        "  ASSET CLASSIFICATION SUMMARY REPORT",
        "=" * 62,
        f"  Total Assets Processed       : {summary['total_assets_processed']}",
        f"  Total Classification Records : {summary['total_classification_records']}",
        f"  Assets Requiring Review      : {summary['manual_review_count']}",
        "",
    ]

    for system, stats in summary["systems"].items():
        lines += [
            f"  [{system}]",
            f"    Assets Classified   : {stats['assets_classified']}",
            f"    Average Top Score   : {stats['average_top_score']}",
            f"    High Confidence     : {stats['high_confidence']}",
            f"    Medium Confidence   : {stats['medium_confidence']}",
            f"    Low Confidence      : {stats['low_confidence']}",
            f"    Flagged for Review  : {stats['flagged_for_review']}",
            "",
        ]

    if summary["assets_requiring_manual_review"]:
        lines.append("  Assets Requiring Manual Review:")
        for aid in summary["assets_requiring_manual_review"]:
            lines.append(f"    - {aid}")
        lines.append("")

    # Cross-system consistency
    if "cross_system_consistency" in summary:
        cs = summary["cross_system_consistency"]
        lines += [
            "  CROSS-SYSTEM CONSISTENCY",
            f"    Consistent assets   : {cs['consistent_count']} / {cs.get('consistent_count',0) + cs.get('inconsistent_count',0)} ({cs['consistency_pct']}%)",
        ]
        if cs["inconsistent_assets"]:
            lines.append("    Inconsistent assets (category disagrees across systems):")
            for aid in cs["inconsistent_assets"]:
                lines.append(f"      - {aid}")
        lines.append("")

    # Data quality
    if "data_quality" in summary:
        dq = summary["data_quality"]
        lines += [
            "  DATA QUALITY",
            f"    Total rows          : {dq['total_rows']}",
            f"    Duplicate asset IDs : {dq['duplicate_count']}",
        ]
        for col, stats in dq.get("fields", {}).items():
            lines.append(f"    {col:<28}: {stats['pct_populated']}% populated")
        if dq.get("duplicate_asset_ids"):
            for aid in dq["duplicate_asset_ids"]:
                lines.append(f"      duplicate: {aid}")
        lines.append("")

    lines.append("=" * 62)

    with open(txt_path, "w") as f:
        f.write("\n".join(lines) + "\n")

    return {"json": str(json_path), "txt": str(txt_path)}
